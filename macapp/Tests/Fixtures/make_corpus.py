"""Writes the UI-test corpus in ``corpus/`` next to this script.

Every file is made up: the places, people and dates exist nowhere else, and each file carries
one invented word (its *token*) that no other file uses, so a test can tell which file a hit
came from. ``README.md`` lists what the tests rely on.

The output is byte-for-byte deterministic, so rerunning this leaves ``git status`` clean:

    garage_python/.venv/bin/python macapp/Tests/Fixtures/make_corpus.py

The text files are written from the strings below, the PDF by hand (one page, the standard
Helvetica font, no compression), and the Word file with python-docx, with fixed document
properties and zip timestamps.
"""

from __future__ import annotations

import io
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path

CORPUS = Path(__file__).resolve().parent / "corpus"

MARKDOWN = """\
# The Quillon Bridge

The Quillon Bridge opened in 1893 and crosses the river Senn at the town of Harrowby.
Its engineer, Adela Morcombe, built it from pale limestone quarried at Fenwick Edge.
Townspeople call the middle arch the zorvexine arch, after the swallows that nest under it.

## Repairs

The bridge was closed for repairs in 1958 after a winter flood cracked the eastern pier.
It reopened the following spring with iron ties set into every arch.
"""

PLAINTEXT = """\
Marrowgate Lighthouse

Idra Voss kept the Marrowgate lighthouse for forty-one years, from 1902 until 1943.
She lit the lamp every evening at dusk and wound the clockwork twice a night.
Her logbook records the plimbrate storm of 1927, when the lantern glass cracked in three places.
The lighthouse was automated in 1961 and now guides ships without a keeper.
"""

CODE = """\
//! Tide tables for Port Ellery, where the tessaroon gauge has read the harbour since 1874.

/// High water at Port Ellery comes 52 minutes later each day.
pub const DAILY_DRIFT_MINUTES: u32 = 52;

/// The time of high water `days` after a high water at `start_minutes` past midnight.
pub fn next_high_water(start_minutes: u32, days: u32) -> u32 {
    (start_minutes + days * DAILY_DRIFT_MINUTES) % (24 * 60)
}
"""

EMAIL = """\
From: Oren Pask <oren.pask@brindlecombe.example>
To: Garage Fixture <fixture@garage.example>
Subject: The Brindlecombe lantern festival
Date: Sat, 04 Oct 2025 09:30:00 +0000
Message-ID: <lantern-festival-2025@brindlecombe.example>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 7bit

Hello,

The Brindlecombe lantern festival is held on the third Saturday of October.
This year the wendleflock parade leaves the village green at seven in the evening.
Every lantern is made of paper and willow, and the tallest one is carried by the miller's family.

See you there,
Oren
"""

PDF_TITLE = "The Ashvale Orchard Survey"
PDF_LINES = [
    "The Ashvale Orchard Survey",
    "",
    "The Ashvale orchard grows 212 varieties of pear on the south slope of Kettle Hill.",
    "The survey was carried out by Nell Oduya in the autumn of 1979.",
    "The oldest tree, known as the orbanquet pear, was planted in 1802.",
    "Most of the fruit is pressed into perry at the Ashvale mill each November.",
]

DOCX_TITLE = "A History of Tavish Glassworks"
DOCX_PARAGRAPHS = [
    "Tavish Glassworks was founded by Mirela Tavish in 1911 in the village of Lowmere.",
    "The works made bottles and window panes until 1934, when it turned to coloured glass.",
    "Its best known piece is the glimmerhaft window in the Lowmere chapel, set in 1938.",
    "The furnace was put out for the last time in 1987.",
]

# Everything python-docx would otherwise stamp with the current time.
FIXED_TIME = datetime(2025, 1, 1, tzinfo=UTC)
ZIP_TIME = (2025, 1, 1, 0, 0, 0)


def _pdf_string(text: str) -> str:
    return "(" + text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def make_pdf() -> bytes:
    """One page of text in the standard Helvetica font, one line per ``T*``."""
    content_lines = ["BT", "/F1 12 Tf", "16 TL", "72 720 Td"]
    for line in PDF_LINES:
        content_lines.append(f"{_pdf_string(line)} Tj T*")
    content_lines.append("ET")
    content = ("\n".join(content_lines) + "\n").encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
        f"<< /Title {_pdf_string(PDF_TITLE)} /Author (Nell Oduya) >>".encode("latin-1"),
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 6 0 R >>\n".encode())
    out.write(f"startxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx() -> bytes:
    """A Word file whose bytes do not depend on when or where it was made."""
    import docx  # python-docx, one of garage_python's dependencies

    document = docx.Document()
    props = document.core_properties
    props.title = DOCX_TITLE
    props.author = "Mirela Tavish"
    props.last_modified_by = "Mirela Tavish"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1
    document.add_heading(DOCX_TITLE, level=1)
    for paragraph in DOCX_PARAGRAPHS:
        document.add_paragraph(paragraph)

    raw = io.BytesIO()
    document.save(raw)
    # Rewrite the zip with fixed entry times, keeping the entry order ([Content_Types].xml first).
    # (Deflate output can differ between zlib builds; any rebuild still extracts to the same text.)
    fixed = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, source.read(info.filename))
    return fixed.getvalue()


FILES = {
    "quillon-bridge.md": lambda: MARKDOWN.encode(),
    "marrowgate-lighthouse.txt": lambda: PLAINTEXT.encode(),
    "tide_tables.rs": lambda: CODE.encode(),
    "lantern-festival.eml": lambda: EMAIL.replace("\n", "\r\n").encode(),
    "ashvale-orchard.pdf": make_pdf,
    "tavish-glassworks.docx": make_docx,
}


def main() -> int:
    CORPUS.mkdir(parents=True, exist_ok=True)
    for name, build in FILES.items():
        (CORPUS / name).write_bytes(build())
        print(f"wrote {CORPUS / name}")
    stray = sorted(p.name for p in CORPUS.iterdir() if p.name not in FILES)
    if stray:
        print(f"not made by this script (delete them or add them here): {', '.join(stray)}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
